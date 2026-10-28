from __future__ import annotations

import io

import pytest
from docx import Document as WordDocument
from httpx import ASGITransport, AsyncClient
from PIL import Image
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import get_settings
from app.db import get_session
from app.main import create_app
from app.models import ImportJob, JobStatus, KnowledgeFile
from app.services.ingestion import TransientImportError, process_import_job
from app.services.ocr import CloudOcr, OcrError
from app.services.parsers import parse_bytes
from app.services.storage import ObjectStorage
from app.tasks import process_import_job_task

ADMIN_HEADERS = {"X-Demo-Role": "admin"}
MANAGER_HEADERS = {"X-Demo-Role": "manager"}


class FailingOcr(CloudOcr):
    is_available = True

    async def extract_markdown(self, image: bytes, mime_type: str) -> str:
        del image, mime_type
        raise OcrError("test OCR outage")


class FailingStorage(ObjectStorage):
    async def put_bytes(self, key: str, data: bytes, content_type: str) -> None:
        del key, data, content_type

    async def get_bytes(self, key: str) -> bytes:
        del key
        raise ConnectionError("temporary storage outage")

    async def delete(self, key: str) -> None:
        del key
        raise ConnectionError("temporary storage outage")


async def create_mock_embedding(client: AsyncClient) -> str:
    response = await client.post(
        "/api/v1/models",
        headers=ADMIN_HEADERS,
        json={
            "display_name": "Reliability embedding",
            "provider": "mock",
            "model_type": "embedding",
            "model_name": "deterministic-test-vector",
        },
    )
    assert response.status_code == 201
    return response.json()["id"]


async def create_mock_ocr(client: AsyncClient) -> str:
    response = await client.post(
        "/api/v1/models",
        headers=ADMIN_HEADERS,
        json={
            "display_name": "Reliability OCR",
            "provider": "mock",
            "model_type": "ocr",
            "model_name": "deterministic-test-ocr",
        },
    )
    assert response.status_code == 201
    return response.json()["id"]


async def create_knowledge_base(client: AsyncClient, *, ocr_model_id: str | None = None) -> str:
    embedding_model_id = await create_mock_embedding(client)
    response = await client.post(
        "/api/v1/knowledge-bases",
        headers=MANAGER_HEADERS,
        json={"name": "Reliability KB", "embedding_model_id": embedding_model_id, "ocr_model_id": ocr_model_id},
    )
    assert response.status_code == 201, response.text
    return response.json()["id"]


def word_payload(*, include_text: bool) -> bytes:
    image = Image.new("RGB", (40, 30), color="white")
    image_payload = io.BytesIO()
    image.save(image_payload, format="PNG")
    image_payload.seek(0)
    document = WordDocument()
    if include_text:
        document.add_paragraph("图片识别失败时仍需保留的正文")
    document.add_picture(image_payload)
    payload = io.BytesIO()
    document.save(payload)
    return payload.getvalue()


async def upload_markdown(client: AsyncClient, knowledge_base_id: str) -> dict:
    response = await client.post(
        f"/api/v1/knowledge-bases/{knowledge_base_id}/files",
        headers=MANAGER_HEADERS,
        files={"upload": ("reliability.md", b"# Reliable\n\nKeep this text.", "text/markdown")},
    )
    assert response.status_code == 202, response.text
    assert response.json()["job"]["status"] == "completed"
    return response.json()


async def test_legacy_ocr_environment_is_used_by_ingestion(
    client: AsyncClient,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    knowledge_base_id = await create_knowledge_base(client)
    monkeypatch.setenv("APP_OCR_PROVIDER", "vision")
    monkeypatch.setenv("APP_OCR_API_KEY", "legacy-test-key")
    get_settings.cache_clear()

    async def extract(self, image: bytes, mime_type: str) -> str:
        del self, image, mime_type
        return "legacy OCR text"

    from app.services.ocr import VisionCloudOcr

    monkeypatch.setattr(VisionCloudOcr, "extract_markdown", extract)
    image = io.BytesIO()
    Image.new("RGB", (40, 30), color="white").save(image, format="PNG")
    response = await client.post(
        f"/api/v1/knowledge-bases/{knowledge_base_id}/files",
        headers=MANAGER_HEADERS,
        files={"upload": ("legacy.png", image.getvalue(), "image/png")},
    )

    assert response.status_code == 202, response.text
    assert response.json()["job"]["status"] == "completed"
    chunks = await client.get(f"/api/v1/knowledge-bases/{knowledge_base_id}/chunks", headers=MANAGER_HEADERS)
    assert "legacy OCR text" in chunks.json()[0]["content"]


async def test_docx_ocr_failure_preserves_text_with_warning_and_scan_fails() -> None:
    parsed = await parse_bytes(word_payload(include_text=True), "mixed.docx", ocr=FailingOcr())

    assert "仍需保留的正文" in parsed.units[0].content
    assert parsed.warnings and "OCR 失败" in parsed.warnings[0]
    assert "OCR 失败" not in parsed.units[0].content

    with pytest.raises(ValueError, match="No readable content"):
        await parse_bytes(word_payload(include_text=False), "scan.docx", ocr=FailingOcr())


async def test_ingestion_persists_ocr_warning_in_metadata_and_stage(
    client: AsyncClient,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    ocr_model_id = await create_mock_ocr(client)
    knowledge_base_id = await create_knowledge_base(client, ocr_model_id=ocr_model_id)
    monkeypatch.setattr("app.services.ingestion.build_ocr", lambda model: FailingOcr())
    response = await client.post(
        f"/api/v1/knowledge-bases/{knowledge_base_id}/files",
        headers=MANAGER_HEADERS,
        files={"upload": ("mixed.docx", word_payload(include_text=True), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )

    assert response.status_code == 202, response.text
    body = response.json()
    assert body["job"]["status"] == "completed"
    assert body["file"]["document_metadata"]["processing_status"] == "warning"
    assert body["file"]["document_metadata"]["warnings"]
    assert any(stage["name"] == "ocr_extraction" and stage["status"] == "warning" for stage in body["job"]["stage_details"])


async def test_terminal_job_is_idempotent_and_transient_failure_can_be_retried(
    client: AsyncClient,
    db_session: AsyncSession,
) -> None:
    knowledge_base_id = await create_knowledge_base(client)
    body = await upload_markdown(client, knowledge_base_id)

    completed = await process_import_job(body["job"]["id"], db_session, storage=FailingStorage())
    assert completed.status == JobStatus.COMPLETED

    file = await db_session.get(KnowledgeFile, body["file"]["id"])
    retry = ImportJob(file_id=file.id, status=JobStatus.QUEUED)
    db_session.add(retry)
    await db_session.commit()
    await db_session.refresh(retry)

    with pytest.raises(TransientImportError):
        await process_import_job(retry.id, db_session, storage=FailingStorage())
    await db_session.refresh(retry)
    await db_session.refresh(file)
    assert retry.status == JobStatus.QUEUED
    assert retry.current_stage == "retry_pending"
    assert file.status == JobStatus.QUEUED

    failed = await process_import_job(retry.id, db_session, storage=FailingStorage(), retry_transient=False)
    assert failed.status == JobStatus.FAILED
    unchanged = await process_import_job(retry.id, db_session, storage=FailingStorage())
    assert unchanged.status == JobStatus.FAILED


def test_celery_autoretries_transient_import_failure(monkeypatch: pytest.MonkeyPatch) -> None:
    calls = 0

    class FakeSession:
        async def __aenter__(self):
            return self

        async def __aexit__(self, *args):
            return None

    class FakeEngine:
        async def dispose(self) -> None:
            return None

    flags: list[bool] = []

    async def fail(job_id: str, session: AsyncSession, *, retry_transient: bool = True) -> ImportJob:
        nonlocal calls
        del job_id, session
        calls += 1
        flags.append(retry_transient)
        raise TransientImportError("temporary connection failure")

    monkeypatch.setattr("app.tasks.build_engine", lambda: FakeEngine())
    monkeypatch.setattr("app.tasks.async_sessionmaker", lambda *args, **kwargs: (lambda: FakeSession()))
    monkeypatch.setattr("app.tasks.process_import_job", fail)

    result = process_import_job_task.apply(args=("job-id",))

    assert result.failed()
    assert calls == 4
    # The final attempt must not stay "retry_pending" forever.
    assert flags == [True, True, True, False]


async def test_celery_dispatch_failure_marks_failed_and_single_file_retry_recovers(
    client: AsyncClient,
    db_session: AsyncSession,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    knowledge_base_id = await create_knowledge_base(client)
    monkeypatch.setenv("APP_TASK_MODE", "celery")
    get_settings.cache_clear()

    def fail_delay(job_id: str) -> None:
        del job_id
        raise ConnectionError("broker unavailable")

    monkeypatch.setattr(process_import_job_task, "delay", fail_delay)
    response = await client.post(
        f"/api/v1/knowledge-bases/{knowledge_base_id}/files",
        headers=MANAGER_HEADERS,
        files={"upload": ("dispatch.md", b"retry me", "text/markdown")},
    )
    assert response.status_code == 503
    assert response.headers["retry-after"] == "5"

    file = (await db_session.scalars(select(KnowledgeFile))).one()
    job = (await db_session.scalars(select(ImportJob))).one()
    assert file.status == JobStatus.FAILED
    assert job.status == JobStatus.FAILED
    assert job.current_stage == "dispatch_failed"

    monkeypatch.setenv("APP_TASK_MODE", "inline")
    get_settings.cache_clear()
    retried = await client.post(
        f"/api/v1/knowledge-bases/{knowledge_base_id}/files/{file.id}/retry",
        headers=MANAGER_HEADERS,
    )
    assert retried.status_code == 202, retried.text
    assert retried.json()["status"] == "completed"


async def test_rebuild_dispatch_failure_does_not_leave_new_job_queued(
    client: AsyncClient,
    db_session: AsyncSession,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    knowledge_base_id = await create_knowledge_base(client)
    await upload_markdown(client, knowledge_base_id)
    monkeypatch.setenv("APP_TASK_MODE", "celery")
    get_settings.cache_clear()
    monkeypatch.setattr(process_import_job_task, "delay", lambda job_id: (_ for _ in ()).throw(ConnectionError()))

    response = await client.post(f"/api/v1/knowledge-bases/{knowledge_base_id}/rebuild", headers=MANAGER_HEADERS)

    assert response.status_code == 503
    jobs = list(await db_session.scalars(select(ImportJob).order_by(ImportJob.created_at)))
    assert jobs[-1].status == JobStatus.FAILED
    assert jobs[-1].current_stage == "dispatch_failed"
    file = (await db_session.scalars(select(KnowledgeFile))).one()
    assert file.status == JobStatus.FAILED


async def test_patch_null_is_rejected_at_route_layer(client: AsyncClient) -> None:
    knowledge_base_id = await create_knowledge_base(client)
    body = await upload_markdown(client, knowledge_base_id)
    chunks = await client.get(f"/api/v1/knowledge-bases/{knowledge_base_id}/chunks", headers=MANAGER_HEADERS)

    kb_response = await client.patch(
        f"/api/v1/knowledge-bases/{knowledge_base_id}",
        headers=MANAGER_HEADERS,
        json={"chunk_size": None},
    )
    chunk_response = await client.patch(
        f"/api/v1/knowledge-bases/{knowledge_base_id}/chunks/{chunks.json()[0]['id']}",
        headers=MANAGER_HEADERS,
        json={"content": None},
    )

    assert kb_response.status_code == 422
    assert chunk_response.status_code == 422
    assert body["job"]["status"] == "completed"


async def test_storage_delete_failure_preserves_db_and_kb_cleanup_removes_objects(
    client: AsyncClient,
    db_session: AsyncSession,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    knowledge_base_id = await create_knowledge_base(client)
    body = await upload_markdown(client, knowledge_base_id)
    file_id = body["file"]["id"]
    file = await db_session.get(KnowledgeFile, file_id)
    storage_key = file.storage_key
    object_path = get_settings().local_storage_path / storage_key
    assert object_path.exists()
    monkeypatch.setattr("app.api.routes.knowledge.get_storage", lambda: FailingStorage())

    file_response = await client.delete(
        f"/api/v1/knowledge-bases/{knowledge_base_id}/files/{file_id}", headers=MANAGER_HEADERS
    )
    kb_response = await client.delete(f"/api/v1/knowledge-bases/{knowledge_base_id}", headers=MANAGER_HEADERS)

    assert file_response.status_code == 503
    assert kb_response.status_code == 503
    assert await db_session.get(KnowledgeFile, file_id) is not None

    from app.services.storage import get_storage

    monkeypatch.setattr("app.api.routes.knowledge.get_storage", get_storage)
    deleted = await client.delete(f"/api/v1/knowledge-bases/{knowledge_base_id}", headers=MANAGER_HEADERS)
    assert deleted.status_code == 204
    assert not object_path.exists()
    assert await db_session.get(KnowledgeFile, file_id) is None


async def test_health_live_ready_and_database_failure(db_session: AsyncSession) -> None:
    class BrokenSession:
        async def execute(self, statement):
            del statement
            raise ConnectionError("database unavailable")

    app = create_app()

    async def broken_session():
        yield BrokenSession()

    app.dependency_overrides[get_session] = broken_session
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as broken_client:
        health = await broken_client.get("/api/v1/health")
        ready = await broken_client.get("/api/v1/health/ready")
        live = await broken_client.get("/api/v1/health/live")

    assert health.status_code == 503
    assert health.json()["database"] == "error"
    assert ready.status_code == 503
    assert ready.json()["database"] == "error"
    assert live.status_code == 200
    assert live.json() == {"status": "ok"}
