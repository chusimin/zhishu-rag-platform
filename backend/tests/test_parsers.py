from __future__ import annotations

import io

import pytest
from docx import Document as WordDocument
from openpyxl import Workbook
from PIL import Image
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.services.ocr import CloudOcr, OcrError
from app.services.parsers import _render_pdf_page, parse_bytes


class StubCloudOcr(CloudOcr):
    is_available = True

    async def extract_markdown(self, image: bytes, mime_type: str) -> str:
        assert image
        assert mime_type.startswith("image/")
        return "流程图 OCR 内容"


def _pdf_with_content(content_bytes: bytes, *, add_blank_page: bool = False) -> bytes:
    writer = PdfWriter()
    first = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    resources = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {NameObject("/F1"): writer._add_object(font)}  # noqa: SLF001 - fixture construction
            )
        }
    )
    content = DecodedStreamObject()
    content.set_data(content_bytes)
    first[NameObject("/Resources")] = resources
    first[NameObject("/Contents")] = writer._add_object(content)  # noqa: SLF001 - fixture construction
    if add_blank_page:
        writer.add_blank_page(width=612, height=792)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


async def test_parse_markdown_preserves_content() -> None:
    parsed = await parse_bytes("# 标题\n\n正文内容".encode(), "guide.md")

    assert parsed.units[0].content == "# 标题\n\n正文内容"
    assert parsed.units[0].metadata["source_type"] == "markdown"


async def test_parse_word_extracts_paragraphs_and_tables() -> None:
    document = WordDocument()
    document.add_heading("产品手册", level=1)
    document.add_paragraph("这是正文。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "字段"
    table.cell(0, 1).text = "说明"
    table.cell(1, 0).text = "状态"
    table.cell(1, 1).text = "已发布"
    payload = io.BytesIO()
    document.save(payload)

    parsed = await parse_bytes(payload.getvalue(), "manual.docx")

    assert parsed.units[0].title == "产品手册"
    assert "这是正文" in parsed.units[0].content
    assert "| 字段 | 说明 |" in parsed.units[0].content
    assert "| 状态 | 已发布 |" in parsed.units[0].content


async def test_word_with_image_requires_cloud_ocr() -> None:
    image = Image.new("RGB", (120, 80), color="white")
    image_payload = io.BytesIO()
    image.save(image_payload, format="PNG")
    image_payload.seek(0)
    document = WordDocument()
    document.add_paragraph("包含业务流程图的文档")
    document.add_picture(image_payload)
    payload = io.BytesIO()
    document.save(payload)

    with pytest.raises(OcrError, match="需要 OCR"):
        await parse_bytes(payload.getvalue(), "with-image.docx")

    parsed = await parse_bytes(payload.getvalue(), "with-image.docx", ocr=StubCloudOcr())
    assert parsed.used_ocr is True
    assert "流程图 OCR 内容" in parsed.units[0].content


@pytest.mark.parametrize(
    ("filename", "payload", "expected"),
    [
        ("records.csv", "姓名,部门\n小明,研发\n".encode(), "| 小明 | 研发 |"),
    ],
)
async def test_parse_tabular_files(filename: str, payload: bytes, expected: str) -> None:
    parsed = await parse_bytes(payload, filename)
    assert expected in parsed.units[0].content


async def test_parse_xlsx_extracts_sheet_as_markdown() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "员工"
    sheet.append(["姓名", "部门"])
    sheet.append(["小明", "研发"])
    payload = io.BytesIO()
    workbook.save(payload)
    workbook.close()

    parsed = await parse_bytes(payload.getvalue(), "employees.xlsx")

    assert parsed.units[0].title == "工作表：员工"
    assert "| 小明 | 研发 |" in parsed.units[0].content


async def test_image_without_cloud_ocr_configuration_fails_explicitly() -> None:
    with pytest.raises(OcrError, match="需要 OCR"):
        await parse_bytes(b"not-a-real-image", "scan.png")


async def test_parse_pdf_extracts_text_and_reports_real_page_count() -> None:
    payload = _pdf_with_content(
        b"BT /F1 12 Tf 72 720 Td (Enterprise knowledge base) Tj ET",
        add_blank_page=True,
    )
    parsed = await parse_bytes(payload, "handbook.pdf")

    assert parsed.page_count == 2
    assert "Enterprise knowledge base" in parsed.units[0].content


async def test_parse_pdf_extracts_ruled_table_as_markdown() -> None:
    drawing = b"\n".join(
        [
            b"0.5 w",
            b"50 700 m 300 700 l S",
            b"50 660 m 300 660 l S",
            b"50 620 m 300 620 l S",
            b"50 620 m 50 700 l S",
            b"170 620 m 170 700 l S",
            b"300 620 m 300 700 l S",
            b"BT /F1 10 Tf 60 678 Td (Name) Tj ET",
            b"BT /F1 10 Tf 180 678 Td (Department) Tj ET",
            b"BT /F1 10 Tf 60 638 Td (Alice) Tj ET",
            b"BT /F1 10 Tf 180 638 Td (Engineering) Tj ET",
        ]
    )

    parsed = await parse_bytes(_pdf_with_content(drawing), "table.pdf")

    assert "[表格 1]" in parsed.units[0].content
    assert "| Name | Department |" in parsed.units[0].content
    assert "| Alice | Engineering |" in parsed.units[0].content


def test_pdf_page_renderer_produces_png_for_cloud_ocr() -> None:
    payload = _pdf_with_content(b"BT /F1 12 Tf 72 720 Td (Scanned page) Tj ET")
    rendered = _render_pdf_page(payload, 0)
    assert rendered.startswith(b"\x89PNG\r\n\x1a\n")
