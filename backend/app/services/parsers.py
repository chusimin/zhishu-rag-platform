from __future__ import annotations

import csv
import io
import mimetypes
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import pdfplumber
from docx import Document as WordDocument
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from openpyxl import load_workbook
from pypdf import PdfReader

from .ocr import CloudOcr, DisabledOcr, OcrError


class UnsupportedFileError(ValueError):
    pass


@dataclass(slots=True)
class ParsedUnit:
    content: str
    title: str = ""
    page_number: int | None = None
    metadata: dict[str, Any] = field(default_factory=dict)


@dataclass(slots=True)
class ParsedDocument:
    units: list[ParsedUnit]
    page_count: int | None = None
    used_ocr: bool = False
    warnings: list[str] = field(default_factory=list)


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".md", ".markdown", ".png", ".jpg", ".jpeg", ".webp", ".xlsx", ".xlsm", ".csv"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}


def _markdown_table(rows: list[list[Any]]) -> str:
    cleaned = [[str(cell).strip() if cell is not None else "" for cell in row] for row in rows]
    cleaned = [row for row in cleaned if any(row)]
    if not cleaned:
        return ""
    width = max(len(row) for row in cleaned)
    normalized = [row + [""] * (width - len(row)) for row in cleaned]
    header = normalized[0]
    separator = ["---"] * width
    return "\n".join(
        ["| " + " | ".join(header) + " |", "| " + " | ".join(separator) + " |"]
        + ["| " + " | ".join(row) + " |" for row in normalized[1:]]
    )


async def _parse_pdf(data: bytes, ocr: CloudOcr) -> ParsedDocument:
    reader = PdfReader(io.BytesIO(data))
    page_count = len(reader.pages)
    units: list[ParsedUnit] = []
    warnings: list[str] = []
    used_ocr = False
    with pdfplumber.open(io.BytesIO(data)) as document:
        for index, page in enumerate(document.pages):
            page_number = index + 1
            parts: list[str] = []
            text = (page.extract_text() or "").strip()
            if text:
                parts.append(text)

            try:
                for table_index, table in enumerate(page.extract_tables()):
                    markdown = _markdown_table(table)
                    if markdown:
                        parts.append(f"[表格 {table_index + 1}]\n{markdown}")
            except Exception:
                # Some malformed PDFs expose text but cannot construct table geometry.
                pass

            image_refs = list(reader.pages[index].images)
            if image_refs and not ocr.is_available:
                raise OcrError(
                    "PDF contains images that require cloud OCR. Select an enabled OCR model for this "
                    "knowledge base; no local OCR model is used."
                )
            if image_refs and ocr.is_available:
                for image_index, image in enumerate(image_refs):
                    try:
                        extension = Path(image.name).suffix.lower().lstrip(".") or "png"
                        ocr_text = (await ocr.extract_markdown(image.data, f"image/{extension}")).strip()
                        if not ocr_text:
                            raise OcrError("Cloud OCR returned empty content")
                        parts.append(f"[图片 {image_index + 1} OCR]\n{ocr_text}")
                        used_ocr = True
                    except Exception as exc:
                        warnings.append(f"PDF 第 {page_number} 页图片 {image_index + 1} OCR 失败：{exc}")

                if not parts:
                    try:
                        rendered = _render_pdf_page(data, index)
                        ocr_text = (await ocr.extract_markdown(rendered, "image/png")).strip()
                        if not ocr_text:
                            raise OcrError("Cloud OCR returned empty content")
                        parts.append(ocr_text)
                        used_ocr = True
                    except Exception as exc:
                        warnings.append(f"PDF 第 {page_number} 页扫描内容 OCR 失败：{exc}")

            if parts:
                units.append(
                    ParsedUnit(
                        content="\n\n".join(parts),
                        title=f"第 {page_number} 页",
                        page_number=page_number,
                        metadata={"source_type": "pdf", "has_images": bool(image_refs)},
                    )
                )
    if not units:
        raise ValueError("No readable content was extracted from the PDF")
    return ParsedDocument(units=units, page_count=page_count, used_ocr=used_ocr, warnings=warnings)


def _render_pdf_page(data: bytes, page_index: int) -> bytes:
    # PDFium is only loaded for an image-only page that actually needs OCR.
    # Text PDFs therefore stay on the pure-Python parsing path.
    import pypdfium2 as pdfium

    document = pdfium.PdfDocument(data)
    try:
        page = document[page_index]
        try:
            bitmap = page.render(scale=2)
            image = bitmap.to_pil()
            output = io.BytesIO()
            image.save(output, format="PNG")
            return output.getvalue()
        finally:
            page.close()
    finally:
        document.close()


async def _parse_word(data: bytes, ocr: CloudOcr) -> ParsedDocument:
    document = WordDocument(io.BytesIO(data))
    units: list[ParsedUnit] = []
    buffer: list[str] = []
    current_title = ""
    used_ocr = False
    warnings: list[str] = []

    def flush() -> None:
        nonlocal buffer
        content = "\n\n".join(item for item in buffer if item.strip()).strip()
        if content:
            units.append(ParsedUnit(content=content, title=current_title, metadata={"source_type": "docx"}))
        buffer = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style and paragraph.style.name.lower().startswith("heading"):
            flush()
            current_title = text
        else:
            buffer.append(text)

    for table_index, table in enumerate(document.tables):
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        markdown = _markdown_table(rows)
        if markdown:
            buffer.append(f"[表格 {table_index + 1}]\n{markdown}")

    image_relationships = [
        relationship
        for relationship in document.part.rels.values()
        if relationship.reltype == RT.IMAGE
    ]
    if image_relationships and not ocr.is_available:
        raise OcrError(
            "Word document contains images that require cloud OCR. Select an enabled OCR model for this "
            "knowledge base; no local OCR model is used."
        )
    if ocr.is_available:
        for image_index, relationship in enumerate(image_relationships, start=1):
            part = relationship.target_part
            mime_type = getattr(part, "content_type", "image/png")
            try:
                text = (await ocr.extract_markdown(part.blob, mime_type)).strip()
                if not text:
                    raise OcrError("Cloud OCR returned empty content")
                buffer.append(f"[文档图片 OCR]\n{text}")
                used_ocr = True
            except Exception as exc:
                warnings.append(f"Word 文档图片 {image_index} OCR 失败：{exc}")
    flush()
    if not units:
        raise ValueError("No readable content was extracted from the Word document")
    return ParsedDocument(units=units, used_ocr=used_ocr, warnings=warnings)


def _parse_markdown(data: bytes) -> ParsedDocument:
    text = data.decode("utf-8-sig").strip()
    if not text:
        raise ValueError("Markdown file is empty")
    return ParsedDocument(units=[ParsedUnit(content=text, title="Markdown", metadata={"source_type": "markdown"})])


async def _parse_image(data: bytes, extension: str, ocr: CloudOcr) -> ParsedDocument:
    mime_type = mimetypes.types_map.get(extension, "image/png")
    text = (await ocr.extract_markdown(data, mime_type)).strip()
    if not text:
        raise OcrError("Cloud OCR returned empty content")
    return ParsedDocument(
        units=[ParsedUnit(content=text, title="图片 OCR", page_number=1, metadata={"source_type": "image"})],
        page_count=1,
        used_ocr=True,
    )


def _parse_workbook(data: bytes) -> ParsedDocument:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    units: list[ParsedUnit] = []
    try:
        for sheet in workbook.worksheets:
            rows = [list(row) for row in sheet.iter_rows(values_only=True)]
            for start in range(0, len(rows), 100):
                batch = rows[start : start + 100]
                markdown = _markdown_table(batch)
                if markdown:
                    units.append(
                        ParsedUnit(
                            content=markdown,
                            title=f"工作表：{sheet.title}",
                            metadata={"source_type": "spreadsheet", "sheet": sheet.title, "row_start": start + 1},
                        )
                    )
    finally:
        workbook.close()
    if not units:
        raise ValueError("Spreadsheet contains no readable cells")
    return ParsedDocument(units=units)


def _parse_csv(data: bytes) -> ParsedDocument:
    text = data.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text)))
    units = []
    for start in range(0, len(rows), 100):
        markdown = _markdown_table(rows[start : start + 100])
        if markdown:
            units.append(
                ParsedUnit(
                    content=markdown,
                    title="CSV 表格",
                    metadata={"source_type": "spreadsheet", "row_start": start + 1},
                )
            )
    if not units:
        raise ValueError("CSV contains no readable rows")
    return ParsedDocument(units=units)


async def parse_bytes(
    data: bytes,
    filename: str,
    *,
    ocr: CloudOcr | None = None,
) -> ParsedDocument:
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileError(f"Unsupported file extension: {extension or 'none'}")
    active_ocr = ocr or DisabledOcr()
    if extension == ".pdf":
        return await _parse_pdf(data, active_ocr)
    if extension == ".docx":
        return await _parse_word(data, active_ocr)
    if extension in {".md", ".markdown"}:
        return _parse_markdown(data)
    if extension in IMAGE_EXTENSIONS:
        return await _parse_image(data, extension, active_ocr)
    if extension in {".xlsx", ".xlsm"}:
        return _parse_workbook(data)
    if extension == ".csv":
        return _parse_csv(data)
    raise UnsupportedFileError(f"Unsupported file extension: {extension}")

